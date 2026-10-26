"""Unit tests for Marp compilation and end-to-end incremental pipeline runner."""

from pathlib import Path

from pptx import Presentation

from distill.compile import generate_marp_deck_markdown, split_slide_budget
from distill.models.curriculum import TopicNode
from distill.models.slide import SlideChunk
from distill.pipeline import DistillPipeline


def test_split_slide_budget_auto_pagination():
    bullets = [
        "First bullet point explaining basic state",
        "Second bullet point describing network latency",
        "Third bullet point on leader heartbeat intervals",
        "Fourth bullet point on candidate term increments",
        "Fifth bullet point on quorum election guarantees",
        "Sixth bullet point on commit index updates",
    ]
    # Max 4 bullets per slide -> should split into 2 slides (4 bullets, then 2 bullets)
    slides = split_slide_budget(
        title="Raft Mechanics",
        bullets=bullets,
        max_bullets_per_slide=4,
        max_chars_per_bullet=50,
        notes="Presenter note for self-testing",
    )

    assert len(slides) == 2
    assert "## Raft Mechanics (1/2)" in slides[0]
    assert "<!-- _notes: Presenter note for self-testing -->" in slides[0]
    assert "## Raft Mechanics (2/2)" in slides[1]


def test_generate_marp_deck_markdown():
    topic = TopicNode(
        id="concurrency",
        title="Locks and Semaphores",
        module_id="mod-01",
        exam_weight=0.20,
    )
    chunk = SlideChunk(
        deck_name="l01.pptx",
        slide_number=1,
        slide_id="l01.pptx#slide-1",
        title="Mutex Invariants",
        text_content="Mutual exclusion guarantees.\nSpinlocks vs Sleep locks.",
    )

    marp_md = generate_marp_deck_markdown(topic, [chunk], chapter_num=1)
    assert "marp: true" in marp_md
    assert "theme: gaia" in marp_md
    assert "# Chapter 1" in marp_md
    assert "Locks and Semaphores" in marp_md
    assert "Mutex Invariants" in marp_md
    assert "Exam Takeaways" in marp_md


def test_pipeline_end_to_end(tmp_path: Path):
    # Setup test workspace
    inbox = tmp_path / "00_inbox_ppts"
    syllabus_dir = tmp_path / "00_syllabus"
    inbox.mkdir(parents=True)
    syllabus_dir.mkdir(parents=True)

    # 1. Create test syllabus
    syllabus_file = syllabus_dir / "syllabus.md"
    syllabus_file.write_text(
        """
# CS6450: Distributed Computing
## Module 1: Foundations
### Logical Clocks
- Objectives: Lamport timestamps
### Distributed Consensus
- Objectives: Paxos and Raft
""",
        encoding="utf-8",
    )

    # 2. Create test presentation
    pptx_file = inbox / "lecture_01.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Logical Clocks and Timestamps"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Lamport clocks establish strict causal ordering."
    prs.save(str(pptx_file))

    # 3. Execute full pipeline
    pipeline = DistillPipeline(root_dir=tmp_path)
    pipeline.run_all(force=True)

    # 4. Verify outputs
    assert (tmp_path / "01_raw_markdown" / "lecture_01.md").exists()
    assert (tmp_path / "02_curriculum_backbone" / "coverage_audit.md").exists()
    assert (tmp_path / "03_study_guide" / "README.md").exists()
    assert (tmp_path / ".cache" / "build_manifest.json").exists()

    # Check that chapters were generated
    study_files = list((tmp_path / "03_study_guide").glob("*.md"))
    assert len(study_files) >= 2  # README.md + at least one chapter

    # Check that Marp decks were generated
    marp_files = list((tmp_path / "04_unified_decks").glob("*.marp.md"))
    assert len(marp_files) >= 1
